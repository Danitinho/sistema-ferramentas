"""
scripts/catalogo.py
===================
Catálogo de produtos — espelho do cadastro do ERP, alimentado por um relatório
exportado em TXT ("todos os produtos").

Para que serve: quando alguém bipa um código de barras no controle de vencidos,
o sistema já sabe o nome do produto, o custo e o preço de venda, em vez de
obrigar a digitar tudo à mão.

O que ele NÃO é: fonte da verdade. É uma FOTO do ERP, com a idade da última
importação. Por isso toda consulta devolve `atualizado_em` — um preço de três
meses atrás não pode ser apresentado como se fosse de hoje.

Importar SUBSTITUI o catálogo inteiro (o arquivo é a foto completa do cadastro):
produto que saiu do ERP some daqui também, e nunca sobra preço fantasma de uma
exportação antiga. A troca é feita numa transação só — se o arquivo falhar no
meio, o catálogo anterior continua inteiro.

Armadilhas do arquivo real (medidas no relatório de 08/09/2026, 95.750 linhas):
  • É UTF-16 LE com BOM, não UTF-8.
  • Os campos vêm entre aspas e separados por '|', mas as aspas NÃO têm escape:
    descrições com polegadas (FACAO M0727N 18") fazem o csv padrão juntar
    colunas e ler o preço errado. Por isso o parse é split('|') + remoção das
    aspas externas, nunca csv.reader.
  • Uma linha traz quebra de linha DENTRO do campo — daí o buffer em
    `_registros`, que acumula até fechar os campos do registro.
  • Preço "R$0,00" é produto sem preço no ERP (11.376 sem venda, 346 sem custo).
    Guardamos NULL: gravar 0,00 num vencido subestimaria a perda no relatório.
  • Códigos com zero à esquerda são cadastros DIFERENTES, com preços diferentes
    ('78924345' R$11,99 × '0000078924345' R$9,39). Por isso a busca é exata —
    ver `consultar`.

Camada de lógica: NÃO importa Flask.
"""
import io
import os
import re
import codecs
import sqlite3
import unicodedata
from datetime import datetime

BASE_DIR  = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DADOS_DIR = os.path.join(BASE_DIR, "dados")
BANCO     = os.path.join(DADOS_DIR, "catalogo.db")

SEP = "|"
LOTE = 5000            # tamanho do bloco de INSERT
MAX_BUFFER = 4000      # trava do acumulador de linha quebrada (caracteres)

# Colunas localizadas PELO NOME no cabeçalho, nunca pela posição: se o layout do
# relatório mudar, a importação falha com mensagem clara em vez de ler a coluna
# errada calada (mesma regra do módulo de reclassificação).
COLUNAS = {
    "codigo_barras":  ("barra1",),
    "descricao":      ("descricao",),
    "venda":          ("precokit",),
    "custo":          ("precocustoaquisicao",),
    "codigo_interno": ("codigo",),          # opcional
}
OBRIGATORIAS = ("codigo_barras", "descricao", "venda", "custo")
ROTULO_COLUNA = {"codigo_barras": "Barra1", "descricao": "Descrição",
                 "venda": "PrecoKit", "custo": "PrecoCustoAquisição"}


class ErroPlanilha(Exception):
    """Arquivo ilegível ou com o cabeçalho fora do contrato."""


# ── Banco ─────────────────────────────────────────────────────────────────────
def _conn():
    os.makedirs(DADOS_DIR, exist_ok=True)
    conn = sqlite3.connect(BANCO)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 5000")
    _init_schema(conn)
    return conn


def _init_schema(conn):
    conn.executescript("""
        -- Espelho do ERP. Sem soft-delete de propósito: a linha não é um fato
        -- histórico, é a foto de agora — quem guarda a história é `importacoes`.
        CREATE TABLE IF NOT EXISTS produtos (
            codigo_barras  TEXT PRIMARY KEY,
            codigo_interno TEXT,
            descricao      TEXT NOT NULL,
            venda          REAL,          -- NULL = sem preço no ERP
            custo          REAL,
            cb_chave       TEXT           -- código sem zeros à esquerda; SÓ para
                                          -- avisar de ambiguidade, nunca para
                                          -- preencher preço (ver `consultar`)
        );
        CREATE INDEX IF NOT EXISTS idx_cat_chave ON produtos(cb_chave);

        -- Trilha de auditoria da importação: quem trocou o catálogo e quando.
        CREATE TABLE IF NOT EXISTS importacoes (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            quando     TEXT NOT NULL,
            usuario    TEXT,
            arquivo    TEXT,
            linhas     INTEGER,
            gravados   INTEGER,
            ignorados  INTEGER,
            conflitos  INTEGER,
            sem_venda  INTEGER,
            sem_custo  INTEGER
        );
    """)
    conn.commit()


def _agora():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _fmt_dt(s):
    """'2026-09-08 14:30:00' → '08/09/2026 14:30'."""
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})', s or "")
    if not m:
        return s or ""
    y, mo, d, h, mi = m.groups()
    return f"{d}/{mo}/{y} {h}:{mi}"


# ── Leitura do arquivo ────────────────────────────────────────────────────────
def _detectar_encoding(cabeca):
    """A exportação de hoje é UTF-16 LE com BOM; as outras formas são rede de
    segurança para o dia em que o ERP mudar de humor."""
    if cabeca.startswith(codecs.BOM_UTF16_LE) or cabeca.startswith(codecs.BOM_UTF16_BE):
        return "utf-16"
    if cabeca.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if len(cabeca) >= 4 and cabeca[1] == 0 and cabeca[3] == 0:
        return "utf-16-le"          # UTF-16 sem BOM
    try:
        cabeca.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1252"


def _abrir(caminho):
    with open(caminho, "rb") as fb:
        enc = _detectar_encoding(fb.read(4))
    return io.open(caminho, encoding=enc, errors="replace", newline="")


def _limpar(campo):
    """Tira espaços e UMA aspa de cada ponta. Aspas internas (as polegadas da
    descrição) ficam onde estão."""
    s = (campo or "").strip()
    if s.startswith('"'):
        s = s[1:]
    if s.endswith('"'):
        s = s[:-1]
    return s.strip()


def _registros(f, ncol):
    """Gera listas de campos já limpos. Acumula linhas até fechar `ncol` campos,
    porque há registro com quebra de linha no meio de um campo."""
    buf = ""
    for bruta in f:
        buf += bruta.rstrip("\r\n")
        if not buf.strip():
            buf = ""
            continue
        partes = buf.split(SEP)
        if len(partes) >= ncol:
            yield [_limpar(p) for p in partes]
            buf = ""
        elif len(buf) > MAX_BUFFER:
            buf = ""    # linha corrompida: descarta em vez de engolir o resto


def _chave_cab(nome):
    """'PrecoCustoAquisição' → 'precocustoaquisicao'."""
    s = unicodedata.normalize("NFKD", _limpar(nome))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _mapear_cabecalho(cab):
    """Posição de cada coluna, pelo nome. Devolve (posicoes, erro)."""
    achadas = {_chave_cab(nome): i for i, nome in enumerate(cab)}
    pos = {}
    for campo, apelidos in COLUNAS.items():
        for a in apelidos:
            if a in achadas:
                pos[campo] = achadas[a]
                break
    faltando = [c for c in OBRIGATORIAS if c not in pos]
    if faltando:
        return None, ("Cabeçalho fora do esperado: falta a coluna "
                      + ", ".join(ROTULO_COLUNA[c] for c in faltando)
                      + ". Colunas encontradas: "
                      + ", ".join(_limpar(c) for c in cab) + ".")
    return pos, None


_SO_DIGITOS = re.compile(r"\D")


def _parse_dinheiro(txt):
    """'R$1.234,56' → 1234.56. Vazio, zero ou negativo → None (sem preço)."""
    s = (txt or "").replace("R$", "").replace(" ", "").strip()
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")   # vírgula decimal, ponto = milhar
    try:
        v = round(float(s), 2)
    except ValueError:
        return None
    return v if v > 0 else None


# ── Importação ────────────────────────────────────────────────────────────────
_SQL_INSERT = ("INSERT OR REPLACE INTO produtos "
               "(codigo_barras, codigo_interno, descricao, venda, custo, cb_chave) "
               "VALUES (?,?,?,?,?,?)")


def importar(caminho, arquivo=None, usuario=None):
    """Substitui o catálogo inteiro pelo conteúdo do arquivo.

    Devolve um resumo com o que entrou e o que foi descartado — a importação
    conta em voz alta, senão um arquivo pela metade passaria por completo."""
    arquivo = arquivo or os.path.basename(caminho)
    f = _abrir(caminho)
    try:
        cabeca = None
        for bruta in f:
            if bruta.strip():
                cabeca = bruta.rstrip("\r\n").split(SEP)
                break
        if not cabeca:
            raise ErroPlanilha("Arquivo vazio.")
        pos, erro = _mapear_cabecalho(cabeca)
        if erro:
            raise ErroPlanilha(erro)
        ncol = max(pos.values()) + 1

        i_cb, i_desc = pos["codigo_barras"], pos["descricao"]
        i_venda, i_custo = pos["venda"], pos["custo"]
        i_cod = pos.get("codigo_interno")

        conn = _conn()
        conn.isolation_level = None            # a transação é conduzida na mão
        try:
            conn.execute("BEGIN IMMEDIATE")
            conn.execute("DELETE FROM produtos")

            linhas = gravados = ignorados = conflitos = sem_venda = sem_custo = 0
            vistos = {}
            exemplos = []
            lote = []
            for campos in _registros(f, ncol):
                linhas += 1
                cb = _SO_DIGITOS.sub("", campos[i_cb]) if i_cb < len(campos) else ""
                desc = campos[i_desc] if i_desc < len(campos) else ""
                if not cb or not desc:
                    ignorados += 1          # sem código ou sem nome não serve
                    continue
                venda = _parse_dinheiro(campos[i_venda]) if i_venda < len(campos) else None
                custo = _parse_dinheiro(campos[i_custo]) if i_custo < len(campos) else None
                if venda is None:
                    sem_venda += 1
                if custo is None:
                    sem_custo += 1
                if cb in vistos:
                    # Código repetido no ERP (9 no arquivo de 08/09/2026, com
                    # descrições e preços diferentes). Vence quem TEM preço de
                    # venda — em 3 dos 9 a segunda linha vem com R$0,00, e ficar
                    # com ela deixaria o produto sem preço à toa. Empate: a
                    # última. O número aparece no resumo da importação.
                    conflitos += 1
                    antes_desc, antes_venda = vistos[cb]
                    if antes_desc != desc and len(exemplos) < 5:
                        exemplos.append(f"{cb}: {antes_desc} / {desc}")
                    if venda is None and antes_venda is not None:
                        continue            # a linha que já está no lote é melhor
                else:
                    gravados += 1
                vistos[cb] = (desc, venda)
                cod_int = campos[i_cod] if (i_cod is not None and i_cod < len(campos)) else None
                lote.append((cb, cod_int, desc, venda, custo, cb.lstrip("0") or "0"))
                if len(lote) >= LOTE:
                    conn.executemany(_SQL_INSERT, lote)
                    lote.clear()
            if lote:
                conn.executemany(_SQL_INSERT, lote)

            if gravados == 0:
                conn.execute("ROLLBACK")
                raise ErroPlanilha(
                    "Nenhum produto válido no arquivo — o catálogo anterior foi mantido.")

            quando = _agora()
            conn.execute(
                "INSERT INTO importacoes (quando, usuario, arquivo, linhas, gravados, "
                "ignorados, conflitos, sem_venda, sem_custo) VALUES (?,?,?,?,?,?,?,?,?)",
                (quando, usuario, arquivo, linhas, gravados, ignorados, conflitos,
                 sem_venda, sem_custo))
            conn.execute("COMMIT")
        except Exception:
            try:
                conn.execute("ROLLBACK")
            except sqlite3.Error:
                pass
            raise
        finally:
            conn.close()
    finally:
        f.close()

    return {
        "ok": True, "arquivo": arquivo, "quando": quando, "quando_fmt": _fmt_dt(quando),
        "linhas": linhas, "gravados": gravados, "ignorados": ignorados,
        "conflitos": conflitos, "conflitos_exemplo": exemplos,
        "sem_venda": sem_venda, "sem_custo": sem_custo,
        "msg": (f"{gravados} produtos no catálogo — {linhas} linhas lidas, "
                f"{ignorados} sem código de barras ou nome, "
                f"{conflitos} códigos repetidos."),
    }


# ── Consulta ──────────────────────────────────────────────────────────────────
def _linha(r, exato):
    return {
        "encontrado": True, "exato": exato, "ambiguo": False,
        "codigo_barras": r["codigo_barras"], "descricao": r["descricao"],
        "venda": r["venda"], "custo": r["custo"],
        # código do produto no ERP: o agente geral confere a linha da nota por ele
        "codigo_interno": r["codigo_interno"] or "",
    }


def consultar(codigo_barras):
    """Produto do catálogo pelo código de barras.

    A busca é EXATA. Zeros à esquerda distinguem cadastros de verdade
    ('78924345' e '0000078924345' são produtos com preços diferentes), então
    normalizar o código encheria o formulário com o preço do cadastro errado.
    Quando o código exato não existe mas há variantes só de zeros:
      • uma única variante → devolve marcada como `exato=False`, e a tela diz
        com qual cadastro casou (a pessoa confere antes de salvar);
      • mais de uma → `ambiguo`, sem preencher nada. Melhor não preencher do
        que preencher o preço de outro produto.
    Devolve sempre um dict, com `atualizado_em` (a idade da foto) junto."""
    conn = _conn()
    try:
        base = _base_consulta(conn)
        base.update(_buscar(conn, codigo_barras))
        return base
    finally:
        conn.close()


MAX_LISTA = 2000


def consultar_lista(codigos):
    """Vários códigos de uma vez (a página /catalogo), com as MESMAS regras de
    `consultar` — exato, variante única de zeros, ambíguo. Uma conexão só.
    Mantém a ordem digitada; código repetido sai uma vez."""
    vistos, lista = set(), []
    for c in codigos or []:
        cb = _SO_DIGITOS.sub("", str(c or ""))
        if cb and cb not in vistos:
            vistos.add(cb)
            lista.append(cb)
    if len(lista) > MAX_LISTA:
        raise ValueError(f"no máximo {MAX_LISTA} códigos por consulta")
    conn = _conn()
    try:
        base = _base_consulta(conn)
        itens = []
        for cb in lista:
            r = {"digitado": cb}
            r.update(_buscar(conn, cb))
            r["margem"] = _margem(r)
            r["observacao"] = _observacao(r)
            itens.append(r)
    finally:
        conn.close()
    return {"itens": itens, "atualizado_em": base["atualizado_em"],
            "atualizado_fmt": base["atualizado_fmt"],
            "encontrados": sum(1 for i in itens if i["encontrado"]),
            "repetidos": sum(1 for c in codigos or [] if _SO_DIGITOS.sub("", str(c or ""))) - len(lista)}


def _margem(i):
    """Margem sobre a venda, em %. None sem custo ou sem venda: um 0% ou 100%
    inventado seria pior que a célula vazia."""
    v, c = i.get("venda"), i.get("custo")
    if v is None or c is None or not v > 0:
        return None
    return round((v - c) / v * 100, 2)


def _observacao(i):
    """O texto da tela e do Excel sai daqui, num lugar só."""
    if not i["encontrado"] and i["ambiguo"]:
        return (f"Código não existe; há {len(i['parecidos'])} cadastros que só diferem "
                "nos zeros à esquerda: " + ", ".join(
                    f"{p['codigo_barras']} ({p['descricao']})" for p in i["parecidos"]))
    if not i["encontrado"]:
        return "Não está no catálogo."
    partes = []
    if not i["exato"]:
        partes.append(f"Casou com o cadastro {i['codigo_barras']} (diferença só nos "
                      "zeros à esquerda) — confira.")
    if i.get("venda") is None:
        partes.append("Sem preço de venda no ERP.")
    if i.get("custo") is None:
        partes.append("Sem custo no ERP.")
    return " ".join(partes)


def gerar_excel_lista(res):
    """Excel da consulta em lista: tabela plana com autofiltro (seção 16 — o
    Excel é para trabalhar, não para apresentar). A linha 1 diz de QUANDO é a
    foto do catálogo; sem ela o arquivo circularia com preço velho parecendo
    de hoje. Devolve BytesIO."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Consulta ao catálogo"
    faltam = len(res["itens"]) - res["encontrados"]
    ws["A1"] = (f"Catálogo — foto do ERP de {res['atualizado_fmt'] or 'data desconhecida'}"
                f" · gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
                f" · {len(res['itens'])} código(s), {res['encontrados']} encontrado(s)"
                + (f", {faltam} não encontrado(s)" if faltam else ""))
    ws["A1"].font = Font(bold=True, size=11, color="1F3A5F")
    cols = [("Código de barras", 18), ("Cód. ERP", 11), ("Produto", 46),
            ("Custo", 12), ("Venda", 12), ("Margem", 10), ("Observação", 60)]
    for n, (titulo, larg) in enumerate(cols, 1):
        c = ws.cell(row=2, column=n, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F3A5F")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(n)].width = larg
    falta = PatternFill("solid", start_color="FCEBEB")
    confere = PatternFill("solid", start_color="FDF3DC")
    for lin, i in enumerate(res["itens"], 3):
        valores = [i["digitado"], i.get("codigo_interno") or "",
                   i.get("descricao") if i["encontrado"] else "",
                   i.get("custo"), i.get("venda"),
                   i["margem"] / 100 if i["margem"] is not None else None,
                   i["observacao"]]
        fill = falta if not i["encontrado"] else (confere if not i["exato"] else None)
        for n, v in enumerate(valores, 1):
            c = ws.cell(row=lin, column=n, value=v)
            if n in (1, 2):
                c.number_format = "@"          # zero à esquerda é dado, não enfeite
            elif n in (4, 5):
                c.number_format = '"R$" #,##0.00'
            elif n == 6:
                c.number_format = "0.0%"
            if fill:
                c.fill = fill
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(ws.max_row, 2)}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _base_consulta(conn):
    base = {"encontrado": False, "exato": False, "ambiguo": False, "parecidos": [],
            "atualizado_em": None, "atualizado_fmt": ""}
    ult = conn.execute("SELECT quando FROM importacoes ORDER BY id DESC LIMIT 1").fetchone()
    if ult:
        base["atualizado_em"] = ult["quando"]
        base["atualizado_fmt"] = _fmt_dt(ult["quando"])
    return base


def _buscar(conn, codigo_barras):
    cb = _SO_DIGITOS.sub("", str(codigo_barras or ""))
    out = {"encontrado": False, "exato": False, "ambiguo": False, "parecidos": []}
    if not cb:
        return out
    r = conn.execute("SELECT * FROM produtos WHERE codigo_barras = ?", (cb,)).fetchone()
    if r:
        out.update(_linha(r, True))
        return out
    rs = conn.execute("SELECT * FROM produtos WHERE cb_chave = ? LIMIT 5",
                      (cb.lstrip("0") or "0",)).fetchall()
    if len(rs) == 1:
        out.update(_linha(rs[0], False))
    elif len(rs) > 1:
        out["ambiguo"] = True
        out["parecidos"] = [{"codigo_barras": x["codigo_barras"],
                             "descricao": x["descricao"]} for x in rs]
    return out


def status():
    """Estado do catálogo, para a tela de atualização."""
    conn = _conn()
    try:
        total = conn.execute("SELECT COUNT(*) FROM produtos").fetchone()[0]
        r = conn.execute("SELECT * FROM importacoes ORDER BY id DESC LIMIT 1").fetchone()
    finally:
        conn.close()
    ultima = None
    if r:
        ultima = dict(r)
        ultima["quando_fmt"] = _fmt_dt(r["quando"])
    return {"total": total, "vazio": total == 0, "ultima": ultima}
